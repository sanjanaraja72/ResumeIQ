import os
import re

import docx
import pdfplumber

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None


def _detect_type(path):
    """File name ah illa, ulla irukkura content ah paathu type kandupidikkum."""
    with open(path, "rb") as f:
        head = f.read(2048)
    if b"%PDF" in head:
        return "pdf"
    if head.startswith(b"PK"):
        return "docx"
    return os.path.splitext(path)[1].lstrip(".").lower()


def _read_pdf(path):
    # 1) pdfplumber
    try:
        with pdfplumber.open(path) as pdf:
            text = "\n".join((p.extract_text() or "") for p in pdf.pages)
        if len(text.strip()) >= 50:
            return text
    except Exception:
        pass

    # 2) PyMuPDF (romba tolerant)
    if fitz is not None:
        try:
            doc = fitz.open(path)
            text = "\n".join(page.get_text() for page in doc)
            doc.close()
            if len(text.strip()) >= 50:
                return text
        except Exception:
            pass

    # 3) pypdf
    if PdfReader is not None:
        try:
            reader = PdfReader(path, strict=False)
            text = "\n".join((p.extract_text() or "") for p in reader.pages)
            if len(text.strip()) >= 50:
                return text
        except Exception:
            pass

    return ""


def extract_text(path):
    kind = _detect_type(path)
    text = ""
    if kind == "pdf":
        text = _read_pdf(path)
    elif kind == "docx":
        d = docx.Document(path)
        text = "\n".join(p.text for p in d.paragraphs)
        for table in d.tables:
            for row in table.rows:
                text += "\n" + " ".join(c.text for c in row.cells)
    else:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
    return text.strip()


def extract_contact(text):
    email = re.search(r"[\w\.-]+@[\w\.-]+\.\w+", text)
    phone = re.search(r"(\+?\d[\d\s\-]{8,}\d)", text)
    linkedin = re.search(r"linkedin\.com/[\w\-/]+", text, re.I)
    github = re.search(r"github\.com/[\w\-/]+", text, re.I)
    first_line = next((l.strip() for l in text.splitlines() if l.strip()), "Candidate")
    return {
        "name": first_line[:40],
        "email": email.group(0) if email else None,
        "phone": phone.group(0).strip() if phone else None,
        "linkedin": linkedin.group(0) if linkedin else None,
        "github": github.group(0) if github else None,
    }